from datetime import date, datetime
from typing import Union
import openpyxl

from common.const import GO_LIVE_DATE

# Khai báo chính xác tên tiêu đề ở Hàng 1
EXACT_START_DATE_HEADER = "StartDate*"
EXACT_BILLING_CONTACT_HEADER = "Billing Contact*"


def clean_sah_funding_episodes(
    workbook: openpyxl.Workbook,
    go_live_date: Union[str, datetime, date] = GO_LIVE_DATE,
    sheet_name: str = "SAHFundingEpisodes",
) -> openpyxl.Workbook:
    """Xử lý cập nhật dữ liệu trên worksheet 'SAHFundingEpisodes':

    1. Gán StartDate* = chuỗi ISO 8601 (YYYY-MM-DD), ví dụ '2025-11-01'.
    2. Thay thế giá trị 'RK_CLIENT' thành 'CLIENT' ở cột Billing Contact*.

    - Hàng 1: Header
    - Hàng 2: Guideline -> Bỏ qua
    - Hàng 3+: Cập nhật dữ liệu
    """
    if sheet_name not in workbook.sheetnames:
        print(
            f"⏩ Worksheet '{sheet_name}' không tồn tại trong workbook -> Bỏ qua."
        )
        return workbook

    sheet = workbook[sheet_name]

    # --- Chuẩn hóa go_live_date thành chuỗi định dạng ISO 8601 (YYYY-MM-DD) ---
    if isinstance(go_live_date, (datetime, date)):
        formatted_date = go_live_date.strftime("%Y-%m-%d")
    else:
        formatted_date = str(go_live_date).strip()

    start_date_col_idx = None
    billing_contact_col_idx = None

    # 1. Tìm vị trí cột ở Hàng 1 (So sánh nguyên văn chuỗi)
    for col in range(1, sheet.max_column + 1):
        cell_val = sheet.cell(row=1, column=col).value
        if cell_val is not None:
            val_str = str(cell_val)
            if val_str == EXACT_START_DATE_HEADER:
                start_date_col_idx = col
            elif val_str == EXACT_BILLING_CONTACT_HEADER:
                billing_contact_col_idx = col

    # 2. Thực hiện cập nhật dữ liệu từ Hàng 3 trở đi
    if start_date_col_idx or billing_contact_col_idx:
        updated_start_date_count = 0
        replaced_billing_contact_count = 0

        for row in range(3, sheet.max_row + 1):
            # Task a: Gán chuỗi ISO 8601 Date cho StartDate*
            if start_date_col_idx:
                sheet.cell(row=row, column=start_date_col_idx).value = (
                    formatted_date
                )
                updated_start_date_count += 1

            # Task b: Thay thế 'RK_CLIENT' -> 'CLIENT' ở Billing Contact*
            if billing_contact_col_idx:
                contact_cell = sheet.cell(
                    row=row, column=billing_contact_col_idx
                )
                if (
                    contact_cell.value is not None
                    and str(contact_cell.value) == "RK_CLIENT"
                ):
                    contact_cell.value = "CLIENT"
                    replaced_billing_contact_count += 1

        print(
            f"✨ Sheet '{sheet_name}': "
            f"Đã gán StartDate* = '{formatted_date}' cho {updated_start_date_count} dòng. "
            f"Đã thay thế {replaced_billing_contact_count} ô 'RK_CLIENT' -> 'CLIENT'."
        )
    else:
        print(
            f"⚠️ Sheet '{sheet_name}': Không tìm thấy cả 2 cột '{EXACT_START_DATE_HEADER}' "
            f"và '{EXACT_BILLING_CONTACT_HEADER}' ở Hàng 1 -> Bỏ qua."
        )

    return workbook