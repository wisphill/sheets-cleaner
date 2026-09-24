import openpyxl

# Tên chính xác của các tiêu đề cột (so sánh nguyên văn)
EXACT_SUPPLIER_CODE_HEADER = "SupplierCode"
EXACT_EMPLOYEE_ID_HEADER = "EmployeeId*"
EXACT_MAIN_ID_HEADER = "Id*"


def remove_employees_with_supplier_code(
    workbook: openpyxl.Workbook,
    employees_sheet_name: str = "Employees",
) -> openpyxl.Workbook:
    """1.

    Tìm các Employee có 'SupplierCode' có giá trị trong sheet 'Employees' và xóa
    dòng.
    2. Tự động quét tất cả các sheet chứa cột 'EmployeeId*' và xóa các dòng liên
    quan.
    """
    if employees_sheet_name not in workbook.sheetnames:
        print(
            f"⏩ Worksheet '{employees_sheet_name}' không tồn tại trong workbook -> Bỏ qua."
        )
        return workbook

    emp_sheet = workbook[employees_sheet_name]

    id_col_idx = None
    supplier_col_idx = None

    # --- BƯỚC 1: Tìm cột Id* và SupplierCode ở sheet Employees ---
    for col in range(1, emp_sheet.max_column + 1):
        cell_val = emp_sheet.cell(row=1, column=col).value
        if cell_val is not None:
            val_str = str(cell_val)
            if val_str == EXACT_MAIN_ID_HEADER:
                id_col_idx = col
            elif val_str == EXACT_SUPPLIER_CODE_HEADER:
                supplier_col_idx = col

    if not id_col_idx or not supplier_col_idx:
        print(
            f"⚠️ Sheet '{employees_sheet_name}': Không tìm thấy cột '{EXACT_MAIN_ID_HEADER}' "
            f"hoặc '{EXACT_SUPPLIER_CODE_HEADER}' ở Hàng 1 -> Bỏ qua."
        )
        return workbook

    # --- BƯỚC 2: Thu thập Id* cần xóa và Xóa dòng trong sheet Employees ---
    removed_emp_ids = set()
    emp_rows_to_delete = []

    # Duyệt từ Hàng 3 trở đi (bỏ qua Hàng 2 Guideline)
    for row in range(3, emp_sheet.max_row + 1):
        supplier_val = emp_sheet.cell(row=row, column=supplier_col_idx).value

        # Nếu SupplierCode có giá trị (không None và không phải chuỗi rỗng)
        if supplier_val is not None and str(supplier_val).strip() != "":
            emp_id = emp_sheet.cell(row=row, column=id_col_idx).value
            if emp_id is not None:
                removed_emp_ids.add(str(emp_id))
                emp_rows_to_delete.append(row)

    if not removed_emp_ids:
        print(
            f"ℹ️ Sheet '{employees_sheet_name}': Không có Employee nào có SupplierCode."
        )
        return workbook

    # Xóa các dòng trong Employees (xóa từ dưới lên)
    for r_idx in sorted(emp_rows_to_delete, reverse=True):
        emp_sheet.delete_rows(r_idx)

    print(
        f"🗑️ Sheet '{employees_sheet_name}': Đã xóa {len(emp_rows_to_delete)} dòng có SupplierCode. "
        f"Danh sách EmployeeId* bị xóa: {removed_emp_ids}"
    )

    # --- BƯỚC 3: Quét TẤT CẢ các Sheet khác để dọn dẹp theo EmployeeId* ---
    for sheet_name in workbook.sheetnames:
        if sheet_name == employees_sheet_name:
            continue

        sheet = workbook[sheet_name]
        emp_id_col_idx = None

        # Tim vị trí cột EmployeeId* ở Hàng 1
        for col in range(1, sheet.max_column + 1):
            cell_val = sheet.cell(row=1, column=col).value
            if (
                cell_val is not None
                and str(cell_val) == EXACT_EMPLOYEE_ID_HEADER
            ):
                emp_id_col_idx = col
                break

        # Nếu sheet này có cột EmployeeId*, tiến hành lọc và xóa dòng
        if emp_id_col_idx:
            child_rows_to_delete = []
            for row in range(3, sheet.max_row + 1):
                cell_val = sheet.cell(row=row, column=emp_id_col_idx).value
                if (
                    cell_val is not None
                    and str(cell_val) in removed_emp_ids
                ):
                    child_rows_to_delete.append(row)

            # Xóa các dòng tìm thấy từ dưới lên
            for r_idx in sorted(child_rows_to_delete, reverse=True):
                sheet.delete_rows(r_idx)

            if child_rows_to_delete:
                print(
                    f"🧹 Sheet liên quan '{sheet_name}': Đã xóa {len(child_rows_to_delete)} dòng chứa EmployeeId* bị loại bỏ."
                )

    return workbook