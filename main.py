import duckdb
import openpyxl

from handler_01_delete_sheets import (
    DEFAULT_SHEETS_TO_DELETE,
    delete_sheets,
)

from handler_02_clean_workbooks import (
    clear_branch_name_column,
)

def load_workbook(input_file: str) -> openpyxl.Workbook:
    """Tải file Excel vào bộ nhớ RAM.

    File gốc trên ổ đĩa KHÔNG bị thay đổi.
    """
    return openpyxl.load_workbook(input_file)


def export_to_new_file(workbook: openpyxl.Workbook, output_file: str):
    """Ghi dữ liệu từ RAM ra một file Excel MỚI hoàn toàn."""
    workbook.save(output_file)
    print(f"📁 Đã xuất thành công ra file MỚI: '{output_file}'")


def load_excel_sheet_duckdb(file_path: str, sheet_name: str):
    """Đọc và truy vấn dữ liệu từ một sheet bằng DuckDB."""
    duckdb.install_extension("spatial")
    duckdb.load_extension("spatial")

    query = f"""
        SELECT * 
        FROM st_read('{file_path}', layer='{sheet_name}')
    """
    return duckdb.query(query).df()


# ==========================================
# QUY TRÌNH THỰC THI
# ==========================================
if __name__ == "__main__":
    input_file = "dovida_staging_2026-09-23_010126_Import_Template_v0.0.53.xlsx"
    output_file = "new_data.xlsx"
    sheet_to_remove = "TempSheet"

    # Bước 1: Đọc file gốc vào RAM (chưa đụng gì file gốc)
    wb = load_workbook(input_file=input_file)
    wb = delete_sheets(wb)
    wb = clear_branch_name_column(workbook=wb)

    # Bước 3: Xuất dữ liệu từ RAM ra FILE MỚI
    export_to_new_file(workbook=wb, output_file=output_file)

    # Bước 4: Dùng DuckDB đọc dữ liệu từ FILE MỚI vừa tạo
    # try:
    #     df = load_excel_sheet_duckdb(
    #         file_path=output_file, sheet_name="DataSheet"
    #     )
    #     print("\n--- Dữ liệu từ file MỚI đọc bằng DuckDB ---")
    #     print(df.head())
    # except Exception as e:
    #     print(f"Error when reading file with DuckDB: {e}")