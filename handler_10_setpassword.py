import os
import random
import openpyxl
from dotenv import load_dotenv

load_dotenv()

EXACT_EMPLOYEES_SHEET = "Employees"


def generate_random_test_id() -> str:
    """Sinh chuỗi Id dạng RK_TEST_ + 9 chữ số ngẫu nhiên."""
    random_9_digits = "".join([str(random.randint(0, 9)) for _ in range(9)])
    return f"RK_TEST_{random_9_digits}"


def process_employee_passwords_or_append(
    workbook: openpyxl.Workbook,
    users_data: list[dict],
    sheet_name: str = EXACT_EMPLOYEES_SHEET,
    env_password_key: str = "TEST_USER_PASSWORD",
) -> openpyxl.Workbook:
    """Kiểm tra danh sách user:

    - Nếu Email đã tồn tại: Chỉ cập nhật Password*.
    - Nếu Email chưa tồn tại: Append user mới vào sheet.
    """
    if sheet_name not in workbook.sheetnames:
        print(f"⚠️ Worksheet '{sheet_name}' không tồn tại trong workbook -> Bỏ qua.")
        return workbook

    sheet = workbook[sheet_name]
    password_val = os.getenv(env_password_key)

    if not password_val:
        print(f"⚠️ Biến môi trường '{env_password_key}' chưa được cấu hình hoặc rỗng.")

    # 1. Bật map tiêu đề cột ở Hàng 1
    header_map = {}
    for col in range(1, sheet.max_column + 1):
        cell_val = sheet.cell(row=1, column=col).value
        if cell_val is not None:
            header_map[str(cell_val).strip()] = col

    email_col_idx = header_map.get("EmailAddress*")
    password_col_idx = header_map.get("Password*")

    if not email_col_idx or not password_col_idx:
        print("⚠️ Không tìm thấy cột 'EmailAddress*' hoặc 'Password*' trong sheet.")
        return workbook

    # 2. Map các email hiện có trong sheet kèm số hàng (row index) tương ứng
    existing_emails = {}
    for row in range(2, sheet.max_row + 1):
        cell_value = sheet.cell(row=row, column=email_col_idx).value
        if cell_value:
            existing_emails[str(cell_value).strip().lower()] = row

    # 3. Duyệt qua từng user cần xử lý
    for user in users_data:
        email = user.get("email", "").strip()
        if not email:
            continue

        email_lower = email.lower()

        # TH1: Email ĐÃ TỒN TẠI -> Chỉ cập nhật Password*
        if email_lower in existing_emails:
            target_row = existing_emails[email_lower]
            sheet.cell(row=target_row, column=password_col_idx, value=password_val)
            print(f"🔄 [Update] Đã cập nhật Password* cho email: {email} (Hàng {target_row})")

        # TH2: Email CHƯA TỒN TẠI -> Append thêm user mới
        else:
            test_id = generate_random_test_id()
            last_name = user.get("last_name", "Nguyen")
            first_name = user.get("first_name", "An")

            row_data = [None] * sheet.max_column

            if "Id*" in header_map:
                row_data[header_map["Id*"] - 1] = test_id
            if "LastName*" in header_map:
                row_data[header_map["LastName*"] - 1] = last_name
            if "FirstName*" in header_map:
                row_data[header_map["FirstName*"] - 1] = first_name
            if "EmailAddress*" in header_map:
                row_data[header_map["EmailAddress*"] - 1] = email
            if "Password*" in header_map:
                row_data[header_map["Password*"] - 1] = password_val
            if "Tenant*" in header_map:
                row_data[header_map["Tenant*"] - 1] = "raykay"
            if "EmploymentType" in header_map:
                row_data[header_map["EmploymentType"] - 1] = "Casual"
            if "Status*" in header_map:
                row_data[header_map["Status*"] - 1] = "Active"
            if "PayrollId" in header_map:
                row_data[header_map["PayrollId"] - 1] = "*Place-Holder"

            sheet.append(row_data)
            print(
                f"➕ [Append] Đã thêm thành công user test mới:\n"
                f"   - Id*: {test_id}\n"
                f"   - Email*: {email}"
            )

    return workbook