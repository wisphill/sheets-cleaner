from datetime import date, datetime
from typing import Union
import openpyxl

# Khai báo chính xác tiêu đề cột (So sánh nguyên văn chuỗi)
EXACT_EPISODE_START_DATE_HEADER = "EpisodeStartDate*"
EXACT_CONTRACT_START_DATE_HEADER = "ContractStartDate*"
EXACT_BILLING_CONTACT_HEADER = "Billing Contact*"
GO_LIVE_DATE = "2026-11-01"


def clean_sah_funding_episode_contracts(
    workbook: openpyxl.Workbook,
    go_live_date: Union[str, datetime, date] = GO_LIVE_DATE,
    sheet_name: str = "SAHFundingEpisodeContracts",
) -> openpyxl.Workbook:
    """Xử lý cập nhật dữ liệu trên worksheet 'SAHFundingEpisodeContracts':

    1. Gán 'EpisodeStartDate*' = chuỗi ISO 8601 (YYYY-MM-DD).
    2. Gán 'ContractStartDate*' = chuỗi ISO 8601 (YYYY-MM-DD).
    3. Thay thế giá trị 'RK_CLIENT' thành 'CLIENT' ở cột 'Billing Contact*'.

    - Hàng 1: Header
    - Hàng 2: Guideline -> Bỏ qua
    - Hàng 3+: Cập nhật dữ liệu
    """
    if sheet_name not in workbook.sheetnames:
        print(
            f"⏩ Worksheet '{sheet_name}' không tồn tại trong workbook -> Bỏ qua."
        )
        return workbook

    sheet = workbook[sheet_name]

    # Chuẩn hóa go_live_date thành chuỗi ISO 8601 (YYYY-MM-DD)
    if isinstance(go_live_date, (datetime, date)):
        formatted_date = go_live_date.strftime("%Y-%m-%d")
    else:
        formatted_date = str(go_live_date)

    episode_start_col_idx = None
    contract_start_col_idx = None
    billing_contact_col_idx = None

    # 1. Tìm vị trí các cột ở Hàng 1
    for col in range(1, sheet.max_column + 1):
        cell_val = sheet.cell(row=1, column=col).value
        if cell_val is not None:
            val_str = str(cell_val)
            if val_str == EXACT_EPISODE_START_DATE_HEADER:
                episode_start_col_idx = col
            elif val_str == EXACT_CONTRACT_START_DATE_HEADER:
                contract_start_col_idx = col
            elif val_str == EXACT_BILLING_CONTACT_HEADER:
                billing_contact_col_idx = col

    # 2. Xử lý dữ liệu từ Hàng 3 trở đi
    if (
        episode_start_col_idx
        or contract_start_col_idx
        or billing_contact_col_idx
    ):
        updated_episode_date_count = 0
        updated_contract_date_count = 0
        replaced_billing_contact_count = 0

        for row in range(3, sheet.max_row + 1):
            # Task a: Gán EpisodeStartDate*
            if episode_start_col_idx:
                sheet.cell(row=row, column=episode_start_col_idx).value = (
                    formatted_date
                )
                updated_episode_date_count += 1

            # Task b: Gán ContractStartDate*
            if contract_start_col_idx:
                sheet.cell(row=row, column=contract_start_col_idx).value = (
                    formatted_date
                )
                updated_contract_date_count += 1

            # Task c: Thay thế 'RK_CLIENT' -> 'CLIENT' ở Billing Contact*
            if billing_contact_col_idx:
                contact_cell = sheet.cell(
                    row=row, column=billing_contact_col_idx
                )
                if (
                    contact_cell.value is not None
                    and str(contact_cell.value) == "RK_CLIENT"
                ):
                    contact_cell.value = "CLIENT"
                    replaced_billing_contact_count += 1

        print(
            f"✨ Sheet '{sheet_name}': "
            f"Đã gán EpisodeStartDate* = '{formatted_date}' cho {updated_episode_date_count} dòng. "
            f"Đã gán ContractStartDate* = '{formatted_date}' cho {updated_contract_date_count} dòng. "
            f"Đã thay thế {replaced_billing_contact_count} ô 'RK_CLIENT' -> 'CLIENT'."
        )
    else:
        print(
            f"⚠️ Sheet '{sheet_name}': Không tìm thấy các cột mục tiêu ở Hàng 1 -> Bỏ qua."
        )

    return workbook