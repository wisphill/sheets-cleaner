"""Compare the Employees / Clients sheets of Raykay_transaction.xlsx against the
source exports in data/comparison and write a diff report (xlsx).

Records are matched by email first, then by any phone number (and by Uid for
employees). The report contains, for each entity:
  - Missing: rows in the source exports not found in Raykay_transaction, and
             whether they exist in the staging export (i.e. removed by main.py)
  - Extra:   rows in Raykay_transaction not found in any source export, with
             status details and the likely reason
  - Matched: matched pairs, with how they were matched and name differences

Usage (from repo root):
    python3 scripts/compare_raykay_transaction.py [--transaction PATH] [--out PATH]
"""

import argparse
import csv
import glob
import os
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from typing import Dict, List, Optional, Set

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from handler_13_client_on_hold import get_on_hold_client_ids  # noqa: E402

COMPARISON_DIR = os.path.join(ROOT, "data", "comparison")

DEFAULT_TRANSACTION = os.path.join(ROOT, "Raykay_transaction.xlsx")
DEFAULT_EMPLOYEE_SOURCES = [
    os.path.join(COMPARISON_DIR, "employee-list Active 13.10.2026.xlsx"),
    os.path.join(COMPARISON_DIR, "employee-KP active 13.10.2026.xlsx"),
]
DEFAULT_CLIENT_SOURCE = os.path.join(COMPARISON_DIR, "client-list 13.10. 3.30pm.csv")
DEFAULT_ON_HOLD = os.path.join(ROOT, "data", "Clients on hold as of 13.10.2026.xlsx")
DEFAULT_STAGING = (sorted(glob.glob(os.path.join(ROOT, "dovida_staging_*.xlsx"))) or [None])[-1]

TRANSACTION_LABEL = "Raykay_transaction"
STAGING_LABEL = "staging"

EMAIL_RE = re.compile(r"[\w.+'-]+@[\w-]+(?:\.[\w-]+)+")
PHONE_RE = re.compile(r"\+?\d[\d\s\-().]{6,}\d")
DNU_RE = re.compile(r"\bdnu", re.IGNORECASE)
# Last 8 digits identify an AU number regardless of +61 / 0 / area code prefix.
PHONE_KEY_DIGITS = 8


@dataclass
class Record:
    source: str
    row: int
    id: str
    first_name: str
    last_name: str
    emails: Set[str] = field(default_factory=set)
    phones: Set[str] = field(default_factory=set)
    uid: str = ""
    raw_emails: str = ""
    raw_phones: str = ""
    data: dict = field(default_factory=dict)

    @property
    def full_name(self) -> str:
        return f"{self.first_name} {self.last_name}".strip()

    @property
    def name_key(self) -> str:
        return normalize_name(self.full_name)

    @property
    def last_name_key(self) -> str:
        return normalize_name(self.last_name.split("(")[0])

    def describe(self) -> str:
        return f"{self.id} {self.full_name} <{self.raw_emails or '-'}>"


def to_str(val) -> str:
    return "" if val is None else str(val).strip()


def normalize_name(name: str) -> str:
    return re.sub(r"[^a-z]", "", name.lower())


def extract_emails(*values) -> Set[str]:
    emails = set()
    for val in values:
        emails.update(m.lower() for m in EMAIL_RE.findall(to_str(val)))
    return emails


def extract_phones(*values) -> Set[str]:
    phones = set()
    for val in values:
        for match in PHONE_RE.findall(to_str(val)):
            digits = re.sub(r"\D", "", match)
            if len(digits) >= PHONE_KEY_DIGITS:
                phones.add(digits[-PHONE_KEY_DIGITS:])
    return phones


def join_raw(*values) -> str:
    return " | ".join(to_str(v) for v in values if to_str(v) not in ("", "0"))


# ==============================================================================
# Loaders
# ==============================================================================
def read_sheet_dicts(path: str, sheet_name: Optional[str] = None, skip_rows: int = 0):
    """Yield (excel_row_number, {header: value}) for each non-empty data row."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    headers = [to_str(h) for h in next(rows)]
    for idx, values in enumerate(rows, start=2):
        if idx < 2 + skip_rows:
            continue
        if all(to_str(v) == "" for v in values):
            continue
        yield idx, dict(zip(headers, values))
    wb.close()


def read_template_sheet(path: str, sheet_name: str):
    """Rows of an import-template sheet (row 2 holds column descriptions)."""
    return read_sheet_dicts(path, sheet_name, skip_rows=1)


def load_template_employees(path: str, label: str = TRANSACTION_LABEL) -> List[Record]:
    records = []
    phone_cols = ["MainPhoneNumber", "PersonalPhoneNumber", "OtherPhoneNumber"]
    for row, d in read_template_sheet(path, "Employees"):
        if not to_str(d.get("Id*")):
            continue
        records.append(Record(
            source=f"{label}/Employees",
            row=row,
            id=to_str(d.get("Id*")),
            first_name=to_str(d.get("FirstName*")),
            last_name=to_str(d.get("LastName*")),
            emails=extract_emails(d.get("EmailAddress*")),
            phones=extract_phones(*(d.get(c) for c in phone_cols)),
            uid=to_str(d.get("uid")),
            raw_emails=join_raw(d.get("EmailAddress*")),
            raw_phones=join_raw(*(d.get(c) for c in phone_cols)),
            data=d,
        ))
    return records


def load_template_clients(path: str, label: str = TRANSACTION_LABEL) -> List[Record]:
    records = []
    phone_cols = ["MainPhoneNumber", "PersonalPhoneNumber", "OtherPhoneNumber"]
    for row, d in read_template_sheet(path, "Clients"):
        if not to_str(d.get("Id*")):
            continue
        records.append(Record(
            source=f"{label}/Clients",
            row=row,
            id=to_str(d.get("Id*")),
            first_name=to_str(d.get("FirstName*")),
            last_name=to_str(d.get("LastName*")),
            emails=extract_emails(d.get("EmailAddress")),
            phones=extract_phones(*(d.get(c) for c in phone_cols)),
            raw_emails=join_raw(d.get("EmailAddress")),
            raw_phones=join_raw(*(d.get(c) for c in phone_cols)),
            data=d,
        ))
    return records


def load_employee_source(path: str) -> List[Record]:
    records = []
    name = os.path.basename(path)
    for row, d in read_sheet_dicts(path):
        records.append(Record(
            source=name,
            row=row,
            id=to_str(d.get("Uid")),
            first_name=to_str(d.get("Name")),
            last_name=to_str(d.get("Last  Name")),
            emails=extract_emails(d.get("Email")),
            phones=extract_phones(d.get("Phone")),
            uid=to_str(d.get("Uid")),
            raw_emails=join_raw(d.get("Email")),
            raw_phones=join_raw(d.get("Phone")),
            data=d,
        ))
    return records


def load_client_source(path: str) -> List[Record]:
    records = []
    name = os.path.basename(path)
    phone_cols = ["phone_main", "phone_personal", "phone_other"]
    with open(path, encoding="cp1252", newline="") as f:
        for row, d in enumerate(csv.DictReader(f), start=2):
            records.append(Record(
                source=name,
                row=row,
                id=to_str(d.get("id")),
                first_name=to_str(d.get("first_name")),
                last_name=to_str(d.get("last_name")),
                emails=extract_emails(d.get("email")),
                phones=extract_phones(*(d.get(c) for c in phone_cols)),
                raw_emails=join_raw(d.get("email")),
                raw_phones=join_raw(*(d.get(c) for c in phone_cols)),
                data=d,
            ))
    return records


def load_client_statuses(path: str) -> Dict[str, List[tuple]]:
    """{ClientId: [(Status, StatusReason, EffectiveDate), ...]} in sheet order."""
    statuses: Dict[str, List[tuple]] = defaultdict(list)
    for _, d in read_template_sheet(path, "ClientStatus"):
        statuses[to_str(d.get("ClientId*"))].append((
            to_str(d.get("Status*")),
            to_str(d.get("StatusReason")),
            to_str(d.get("EffectiveDate*")),
        ))
    return statuses


def load_client_services(path: str) -> Dict[str, Counter]:
    """{ClientId: Counter((Status, StatusReason))}"""
    services: Dict[str, Counter] = defaultdict(Counter)
    for _, d in read_template_sheet(path, "Services"):
        services[to_str(d.get("ClientId*"))][
            (to_str(d.get("Status*")), to_str(d.get("StatusReason")))
        ] += 1
    return services


def load_employee_roles(path: str) -> Dict[str, List[str]]:
    roles: Dict[str, List[str]] = defaultdict(list)
    for _, d in read_template_sheet(path, "EmployeeRoles"):
        roles[to_str(d.get("EmployeeId*"))].append(to_str(d.get("Role*")))
    return roles


# ==============================================================================
# Matching
# ==============================================================================
@dataclass
class MatchResult:
    matched: List[tuple]  # (source_record, target_record, method)
    missing: List[Record]
    extra: List[Record]


def build_index(records: List[Record], key_fn) -> Dict[str, List[Record]]:
    index: Dict[str, List[Record]] = {}
    for rec in records:
        for key in key_fn(rec):
            index.setdefault(key, []).append(rec)
    return index


def compare(sources: List[Record], targets: List[Record]) -> MatchResult:
    strategies = [
        ("email", lambda r: r.emails),
        ("phone", lambda r: r.phones),
        ("uid", lambda r: {r.uid} if r.uid else set()),
    ]
    indexes = [(name, fn, build_index(targets, fn)) for name, fn in strategies]

    matched, missing = [], []
    matched_target_ids: Set[int] = set()

    for src in sources:
        found = None
        for method, key_fn, index in indexes:
            candidates = [t for key in key_fn(src) for t in index.get(key, [])]
            if not candidates:
                continue
            # Prefer a candidate with the same name, e.g. when a phone is shared.
            same_name = [t for t in candidates if t.name_key == src.name_key]
            unused = [t for t in (same_name or candidates) if id(t) not in matched_target_ids]
            found = (unused or same_name or candidates)[0]
            matched.append((src, found, method))
            matched_target_ids.add(id(found))
            break
        if found is None:
            missing.append(src)

    extra = [t for t in targets if id(t) not in matched_target_ids]
    return MatchResult(matched, missing, extra)


def name_hints(records: List[Record], others: List[Record]) -> Dict[int, str]:
    """For unmatched records, list records on the other side with the same name."""
    by_name = build_index(others, lambda r: {r.name_key} if r.name_key else set())
    return {
        id(rec): "; ".join(h.describe() for h in by_name.get(rec.name_key, []))
        for rec in records
    }


# ==============================================================================
# Analysis
# ==============================================================================
def staging_lookup(missing: List[Record], staging: Optional[List[Record]]) -> Dict[int, str]:
    """Explain missing rows: present in staging (removed by main.py) or not exported."""
    if staging is None:
        return {id(r): "" for r in missing}
    result = compare(missing, staging)
    notes = {
        id(src): f"In staging as {tgt.describe()} (matched by {method}) -> removed by main.py"
        for src, tgt, method in result.matched
    }
    for rec in result.missing:
        notes[id(rec)] = "Not in staging export -> likely created after the staging export"
    return notes


def analyze_extra_employees(result: MatchResult, targets: List[Record], transaction_path: str):
    """Return (headers, {id(record): cells}) describing why each extra employee is extra."""
    roles = load_employee_roles(transaction_path)
    matched_ids = {id(t) for _, t, _ in result.matched}
    # Phones shared by more than 2 employees are office numbers, not personal ones.
    by_phone = {p: rs for p, rs in build_index(targets, lambda r: r.phones).items() if len(rs) <= 2}
    by_last_name = build_index(targets, lambda r: {r.last_name_key} if r.last_name_key else set())

    details = {}
    for rec in result.extra:
        status = to_str(rec.data.get("Status*"))
        reasons = []

        if rec.id.startswith("RK_TEST_"):
            reasons.append("Test account added by main.py")
        if DNU_RE.search(f"{rec.full_name} {rec.raw_emails}"):
            reasons.append("Marked DNU (do not use)")

        others = {id(t): t for p in rec.phones for t in by_phone.get(p, []) if t is not rec}
        others.update({
            id(t): t for t in by_last_name.get(rec.last_name_key, [])
            if t is not rec and t.first_name[:1].lower() == rec.first_name[:1].lower()
        })
        if others:
            reasons.append("Possible duplicate of " + "; ".join(
                f"{t.describe()} [{'matched' if id(t) in matched_ids else 'extra'}]"
                for t in others.values()
            ))

        if status.lower() != "active":
            reasons.append(f"Status is '{status}' (source lists only contain Active)")
        if not reasons:
            reasons.append("Active in transaction but not in source lists")

        details[id(rec)] = [
            status,
            to_str(rec.data.get("JobTitle")),
            ", ".join(roles.get(rec.id, [])),
            "\n".join(reasons),
        ]
    return ["Status*", "JobTitle", "Roles", "Likely reason"], details


def format_services(counter: Counter) -> str:
    return ", ".join(
        f"{status}{f' ({reason})' if reason else ''} x{n}"
        for (status, reason), n in sorted(counter.items())
    )


def analyze_extra_clients(
    result: MatchResult,
    transaction_path: str,
    on_hold_path: Optional[str],
    staging_path: Optional[str],
):
    """Return (headers, {id(record): cells}) describing why each extra client is extra."""
    statuses = load_client_statuses(transaction_path)
    services = load_client_services(transaction_path)
    staging_statuses = load_client_statuses(staging_path) if staging_path else {}
    on_hold = {}
    if on_hold_path and os.path.exists(on_hold_path):
        on_hold = {f"RK_{k}": v for k, v in get_on_hold_client_ids(on_hold_path).items()}

    details = {}
    for rec in result.extra:
        status, reason, _ = (statuses.get(rec.id) or [("", "", "")])[-1]
        staging_status = (staging_statuses.get(rec.id) or [("", "", "")])[-1]
        svc = services.get(rec.id, Counter())
        svc_statuses = {s for s, _ in svc}
        reasons = []

        if not status:
            reasons.append(
                "No ClientStatus row in transaction"
                + (f" (staging last status: {staging_status[0]} {staging_status[2]})"
                   if staging_status[0] else "")
            )
        if staging_status[0] == "Pending" and not status:
            reasons.append("Pending (new) client - its ClientStatus row was dropped by main.py")
        if status == "On Hold":
            reasons.append(
                "On Hold (source list only contains active clients)"
                + ("" if rec.id in on_hold else " - NOT in the On Hold file")
            )
            if "Active" in svc_statuses:
                reasons.append("Inconsistent: client On Hold but has Active services")
        elif status and status != "Active":
            reasons.append(f"Status is '{status}'")
        if svc and svc_statuses == {"Discharged"}:
            reasons.append(
                "All services Discharged"
                + (" - client likely should be Discharged" if status in ("", "Active", "On Hold") else "")
            )
        if any(r.lower() == "deceased" for _, r in svc):
            reasons.append("A service is Discharged as Deceased")
        if not svc:
            reasons.append("No services")
        if not reasons:
            reasons.append("Active in transaction but not in source list")

        details[id(rec)] = [
            status,
            reason,
            "Yes" if rec.id in on_hold else "",
            to_str(rec.data.get("client_type")),
            format_services(svc),
            "\n".join(reasons),
        ]
    headers = ["ClientStatus", "StatusReason", "In On Hold file", "client_type", "Services", "Likely reason"]
    return headers, details


# ==============================================================================
# Report
# ==============================================================================
HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
WRAP = Alignment(wrap_text=True, vertical="top")


def write_table(ws, headers: List[str], rows: List[list], highlight_col: Optional[int] = None):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for r in rows:
        ws.append(r)
        for cell in ws[ws.max_row]:
            cell.alignment = WRAP
            if highlight_col is not None and r[highlight_col]:
                cell.fill = WARN_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        width = max(len(line) for c in col for line in to_str(c.value).split("\n") or [""])
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 70)


UNMATCHED_HEADERS = ["Source", "Row", "Id", "FirstName", "LastName", "Email", "Phones"]


def unmatched_row(r: Record) -> list:
    return [r.source, r.row, r.id, r.first_name, r.last_name, r.raw_emails, r.raw_phones]


def add_entity_sheets(wb, label, sources, targets, result, extra_analysis, staging):
    hints = name_hints(result.missing, targets)
    staging_notes = staging_lookup(result.missing, staging)
    ws = wb.create_sheet(f"{label} Missing")
    write_table(
        ws,
        UNMATCHED_HEADERS + ["Same name in transaction (possible match)", "In staging export?"],
        [unmatched_row(r) + [hints[id(r)], staging_notes[id(r)]] for r in result.missing],
        highlight_col=7,
    )

    extra_headers, extra_details = extra_analysis
    hints = name_hints(result.extra, sources)
    ws = wb.create_sheet(f"{label} Extra")
    write_table(
        ws,
        UNMATCHED_HEADERS + extra_headers + ["Same name in source lists (possible match)"],
        [unmatched_row(r) + extra_details[id(r)] + [hints[id(r)]] for r in result.extra],
        highlight_col=len(UNMATCHED_HEADERS) + len(extra_headers),
    )

    ws = wb.create_sheet(f"{label} Matched")
    rows = []
    for src, tgt, method in result.matched:
        name_diff = "YES" if src.name_key != tgt.name_key else ""
        rows.append([
            method, name_diff,
            src.source, src.id, src.full_name, src.raw_emails, src.raw_phones,
            tgt.id, tgt.full_name, tgt.raw_emails, tgt.raw_phones,
        ])
    write_table(ws, [
        "Matched by", "Name differs",
        "Source", "Source Id", "Source Name", "Source Email", "Source Phones",
        "Transaction Id", "Transaction Name", "Transaction Email", "Transaction Phones",
    ], rows, highlight_col=1)


def summary_rows(label, sources, targets, result, extra_details, status_col) -> List[list]:
    by_method = Counter(method for _, _, method in result.matched)
    extra_by_status = Counter(extra_details[id(r)][status_col] or "(none)" for r in result.extra)
    name_diffs = sum(1 for s, t, _ in result.matched if s.name_key != t.name_key)
    return [
        [label, "Source rows", len(sources)],
        [label, "Transaction rows", len(targets)],
        [label, "Matched", len(result.matched)],
        *[[label, f"  matched by {m}", n] for m, n in by_method.items()],
        [label, "Matched but name differs", name_diffs],
        [label, "Missing from transaction", len(result.missing)],
        [label, "Extra in transaction (not in sources)", len(result.extra)],
        *[[label, f"  extra with status {s}", n] for s, n in extra_by_status.most_common()],
    ]


def print_unmatched(title: str, records: List[Record], details: Dict[int, list]):
    print(f"\n=== {title} ({len(records)}) ===")
    for r in records:
        print(f"- {r.describe()}")
        for line in details[id(r)][-1].split("\n"):
            if line:
                print(f"    {line}")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--transaction", default=DEFAULT_TRANSACTION)
    parser.add_argument("--employees", nargs="+", default=DEFAULT_EMPLOYEE_SOURCES)
    parser.add_argument("--clients", default=DEFAULT_CLIENT_SOURCE)
    parser.add_argument("--on-hold", default=DEFAULT_ON_HOLD)
    parser.add_argument("--staging", default=DEFAULT_STAGING,
                        help="Staging import template that main.py reads (used to explain diffs)")
    parser.add_argument("--out", default=os.path.join(
        COMPARISON_DIR, f"comparison_report_{datetime.now():%Y-%m-%d_%H%M%S}.xlsx"))
    args = parser.parse_args()
    staging = args.staging if args.staging and os.path.exists(args.staging) else None

    emp_sources = [r for p in args.employees for r in load_employee_source(p)]
    emp_targets = load_template_employees(args.transaction)
    emp_result = compare(emp_sources, emp_targets)
    emp_analysis = analyze_extra_employees(emp_result, emp_targets, args.transaction)
    emp_staging = load_template_employees(staging, STAGING_LABEL) if staging else None

    cli_sources = load_client_source(args.clients)
    cli_targets = load_template_clients(args.transaction)
    cli_result = compare(cli_sources, cli_targets)
    cli_analysis = analyze_extra_clients(cli_result, args.transaction, args.on_hold, staging)
    cli_staging = load_template_clients(staging, STAGING_LABEL) if staging else None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    summary = (summary_rows("Employees", emp_sources, emp_targets, emp_result, emp_analysis[1], 0)
               + summary_rows("Clients", cli_sources, cli_targets, cli_result, cli_analysis[1], 0))
    write_table(ws, ["Entity", "Metric", "Count"], summary)

    add_entity_sheets(wb, "Employees", emp_sources, emp_targets, emp_result, emp_analysis, emp_staging)
    add_entity_sheets(wb, "Clients", cli_sources, cli_targets, cli_result, cli_analysis, cli_staging)

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    wb.save(args.out)

    for entity, metric, count in summary:
        print(f"{entity:<10} {metric:<40} {count}")

    for label, result, staging_records in [
        ("Employees", emp_result, emp_staging), ("Clients", cli_result, cli_staging)
    ]:
        notes = staging_lookup(result.missing, staging_records)
        print_unmatched(f"{label} missing", result.missing, {k: [v] for k, v in notes.items()})
    print_unmatched("Employees extra", emp_result.extra, emp_analysis[1])
    print_unmatched("Clients extra", cli_result.extra, cli_analysis[1])

    print(f"\n✅ Report written to {args.out}")


if __name__ == "__main__":
    main()
