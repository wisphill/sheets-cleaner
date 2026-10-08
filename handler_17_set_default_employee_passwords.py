import os
import sys
import openpyxl
from dotenv import load_dotenv

load_dotenv()

EXACT_EMPLOYEES_SHEET = "Employees"
EXACT_ID_HEADER = "Id*"
EXACT_PASSWORD_HEADER = "Password*"

DEFAULT_PASSWORD_ENV_KEY = "DEFAULT_EMPLOYEE_PASSWORD"

# Theo guideline của cột Password* trong template
PASSWORD_MIN_LENGTH = 8
PASSWORD_MAX_LENGTH = 72


def is_empty_or_none(val) -> bool:
    """Helper kiểm tra giá trị ô có rỗng hay không."""
    return val is None or str(val).strip() == ""


def set_default_employee_passwords(
    workbook: openpyxl.Workbook,
    sheet_name: str = EXACT_EMPLOYEES_SHEET,
    env_password_key: str = DEFAULT_PASSWORD_ENV_KEY,
    overwrite_existing: bool = False,
) -> openpyxl.Workbook:
    """Gán password mặc định (lấy từ biến môi trường) cho mọi Employee.

    - Mặc định chỉ điền vào các ô Password* đang trống, để giữ nguyên password
      đã gán trước đó (VD: user test / admin từ handler_10).
    - overwrite_existing=True: ghi đè password cho toàn bộ Employee.
    """
    if sheet_name not in workbook.sheetnames:
        print(f"⚠️ Worksheet '{sheet_name}' không tồn tại trong workbook -> Bỏ qua.")
        return workbook

    password_val = os.getenv(env_password_key)

    if not password_val:
        print(
            f"❌ Biến môi trường '{env_password_key}' chưa được cấu hình hoặc"
            " rỗng. Dừng chương trình."
        )
        sys.exit(1)

    if not PASSWORD_MIN_LENGTH <= len(password_val) <= PASSWORD_MAX_LENGTH:
        print(
            f"❌ '{env_password_key}' phải có độ dài từ {PASSWORD_MIN_LENGTH} đến"
            f" {PASSWORD_MAX_LENGTH} ký tự. Dừng chương trình."
        )
        sys.exit(1)

    sheet = workbook[sheet_name]

    header_map = {}
    for col in range(1, sheet.max_column + 1):
        cell_val = sheet.cell(row=1, column=col).value
        if cell_val is not None:
            header_map[str(cell_val).strip()] = col

    id_col_idx = header_map.get(EXACT_ID_HEADER)
    password_col_idx = header_map.get(EXACT_PASSWORD_HEADER)

    if not id_col_idx or not password_col_idx:
        print(
            f"⚠️ Không tìm thấy cột '{EXACT_ID_HEADER}' hoặc"
            f" '{EXACT_PASSWORD_HEADER}' trong sheet '{sheet_name}' -> Bỏ qua."
        )
        return workbook

    updated_count = 0
    kept_count = 0

    # Hàng 2 là guideline -> dữ liệu bắt đầu từ Hàng 3
    for row in range(3, sheet.max_row + 1):
        if is_empty_or_none(sheet.cell(row=row, column=id_col_idx).value):
            continue

        password_cell = sheet.cell(row=row, column=password_col_idx)
        if not overwrite_existing and not is_empty_or_none(password_cell.value):
            kept_count += 1
            continue

        password_cell.value = password_val
        updated_count += 1

    print(
        f"🔑 Sheet '{sheet_name}': Đã gán password mặc định cho {updated_count}"
        f" Employee (giữ nguyên {kept_count} Employee đã có password)."
    )

    return workbook
