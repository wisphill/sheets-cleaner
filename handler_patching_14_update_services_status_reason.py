import openpyxl

SERVICES_SHEET_NAME = "Services"
EXACT_STATUS_REASON_HEADER = "StatusReason"

# Current value -> value configured in AlayaCare.
# Key được so khớp không phân biệt hoa/thường nên giá trị đã đúng sẵn vẫn khớp.
STATUS_REASON_MAPPING = {
    "ADHOC/SHORT TERM": "Adhoc/Short Term",
    "BROKER CLIENT": "Broker Client",
    "DECEASED": "Deceased",
    # Sheet ghi 'CARER', bảng cấu hình ghi 'CAREER' -> chấp nhận cả hai.
    "FAMILY/PRIVATE CARER": "Family/Private Carer",
    "FAMILY/PRIVATE CAREER": "Family/Private Carer",
    "HOSPITAL": "Hospital",
    "MOVED TO COMPETITOR": "Moved To Competitor",
    "OTHER": "Other",
    "OTHER ONHOLD": "Other OnHold",
    "RESIDENTIAL CARE": "Residential Care",
    "RESPITE": "Respite",
}

_NORMALISED_MAPPING = {
    key.strip().lower(): value for key, value in STATUS_REASON_MAPPING.items()
}


def handler_services_table_update_status_reason(
    workbook: openpyxl.Workbook,
    sheet_name: str = SERVICES_SHEET_NAME,
) -> openpyxl.Workbook:
    """Chuẩn hóa cột 'StatusReason' của sheet 'Services' về đúng giá trị đã cấu
    hình trong AlayaCare.

    - Hàng 1: Header, Hàng 2: Guideline -> bỏ qua.
    - Giá trị không nằm trong bảng mapping được giữ nguyên và liệt kê ở cuối để
      rà soát thủ công.
    """
    if sheet_name not in workbook.sheetnames:
        print(
            f"⏩ Worksheet '{sheet_name}' does not exist in workbook -> Skipped."
        )
        return workbook

    sheet = workbook[sheet_name]

    status_reason_col_idx = None
    for col in range(1, sheet.max_column + 1):
        cell_val = sheet.cell(row=1, column=col).value
        if cell_val is not None and str(cell_val).strip() == EXACT_STATUS_REASON_HEADER:
            status_reason_col_idx = col
            break

    if not status_reason_col_idx:
        print(
            f"⚠️ Sheet '{sheet_name}': Could not find column"
            f" '{EXACT_STATUS_REASON_HEADER}' on Row 1 -> Skipped."
        )
        return workbook

    updated_per_value = {}
    unmapped_values = {}

    for row in range(3, sheet.max_row + 1):
        cell = sheet.cell(row=row, column=status_reason_col_idx)
        if cell.value is None:
            continue

        current_val = str(cell.value).strip()
        if current_val == "":
            continue

        configured_val = _NORMALISED_MAPPING.get(current_val.lower())

        if configured_val is None:
            unmapped_values[current_val] = unmapped_values.get(current_val, 0) + 1
            continue

        if cell.value != configured_val:
            cell.value = configured_val
            key = f"{current_val} -> {configured_val}"
            updated_per_value[key] = updated_per_value.get(key, 0) + 1

    total_updated = sum(updated_per_value.values())
    print(
        f"✨ Sheet '{sheet_name}': Normalised '{EXACT_STATUS_REASON_HEADER}' for"
        f" {total_updated} row(s)."
    )
    for key, count in sorted(updated_per_value.items()):
        print(f"   • {key}: {count}")

    if unmapped_values:
        total_unmapped = sum(unmapped_values.values())
        print(
            f"⚠️ Sheet '{sheet_name}': {total_unmapped} row(s) have a"
            " StatusReason that is not in the mapping -> left unchanged:"
        )
        for value, count in sorted(
            unmapped_values.items(), key=lambda kv: -kv[1]
        ):
            print(f"   • {value!r}: {count}")

    return workbook
