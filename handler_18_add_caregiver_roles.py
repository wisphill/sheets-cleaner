import openpyxl

from handler_04_add_employee_roles import (
    EXACT_EMP_ID_HEADER,
    EXACT_MAIN_ID_HEADER,
    EXACT_ROLE_HEADER,
    EXACT_TENANT_HEADER,
)

DEFAULT_ROLE = "CAREGiver"
DEFAULT_TENANT = "raykay"


def is_empty_or_none(val) -> bool:
    """Helper kiểm tra giá trị ô có rỗng hay không."""
    return val is None or str(val).strip() == ""


def get_header_map(sheet) -> dict:
    header_map = {}
    for col in range(1, sheet.max_column + 1):
        cell_val = sheet.cell(row=1, column=col).value
        if cell_val is not None:
            header_map[str(cell_val).strip()] = col
    return header_map


def add_caregiver_role_to_all_employees(
    workbook: openpyxl.Workbook,
    role: str = DEFAULT_ROLE,
    tenant: str = DEFAULT_TENANT,
    sheet_name: str = "EmployeeRoles",
    employees_sheet_name: str = "Employees",
) -> openpyxl.Workbook:
    """Thêm Role (mặc định "CAREGiver") vào sheet EmployeeRoles cho mọi Employee.

    EmployeeId* lấy từ cột 'Id*' của sheet Employees. Các Employee đã có Role
    trong EmployeeRoles (VD: System Administrator từ handler_04) được giữ
    nguyên và không bị thêm Role này.
    """
    if employees_sheet_name not in workbook.sheetnames:
        print(
            f"⚠️ Worksheet '{employees_sheet_name}' không tồn tại trong workbook"
            " -> Bỏ qua."
        )
        return workbook

    if sheet_name not in workbook.sheetnames:
        print(
            f"⚠️ Worksheet '{sheet_name}' không tồn tại (cần chạy handler_04"
            " trước) -> Bỏ qua."
        )
        return workbook

    emp_sheet = workbook[employees_sheet_name]
    roles_sheet = workbook[sheet_name]

    emp_id_col = get_header_map(emp_sheet).get(EXACT_MAIN_ID_HEADER)
    if not emp_id_col:
        print(
            f"⚠️ Không tìm thấy cột '{EXACT_MAIN_ID_HEADER}' trong sheet"
            f" '{employees_sheet_name}' -> Bỏ qua."
        )
        return workbook

    roles_header_map = get_header_map(roles_sheet)
    role_emp_id_col = roles_header_map.get(EXACT_EMP_ID_HEADER)
    role_col = roles_header_map.get(EXACT_ROLE_HEADER)
    if not role_emp_id_col or not role_col:
        print(
            f"⚠️ Không tìm thấy cột '{EXACT_EMP_ID_HEADER}' hoặc"
            f" '{EXACT_ROLE_HEADER}' trong sheet '{sheet_name}' -> Bỏ qua."
        )
        return workbook

    # Hàng 2 là guideline -> dữ liệu bắt đầu từ Hàng 3
    employees_with_role = set()
    for row in range(3, roles_sheet.max_row + 1):
        val = roles_sheet.cell(row=row, column=role_emp_id_col).value
        if not is_empty_or_none(val):
            employees_with_role.add(str(val).strip())

    added_count = 0
    seen_ids = set()

    for row in range(3, emp_sheet.max_row + 1):
        val = emp_sheet.cell(row=row, column=emp_id_col).value
        if is_empty_or_none(val):
            continue

        emp_id = str(val).strip()
        if emp_id in employees_with_role or emp_id in seen_ids:
            continue
        seen_ids.add(emp_id)

        row_data = [None] * roles_sheet.max_column
        row_data[role_emp_id_col - 1] = emp_id
        row_data[role_col - 1] = role
        if EXACT_TENANT_HEADER in roles_header_map:
            row_data[roles_header_map[EXACT_TENANT_HEADER] - 1] = tenant

        roles_sheet.append(row_data)
        added_count += 1

    print(
        f"✨ Sheet '{sheet_name}': Đã thêm Role '{role}' cho {added_count}"
        f" Employee (giữ nguyên {len(employees_with_role)} Employee đã có Role)."
    )

    return workbook
