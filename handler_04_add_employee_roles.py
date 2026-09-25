from typing import Dict, List, Optional
import openpyxl
from openpyxl.styles import Font, Alignment

EXACT_EMP_ID_HEADER = "EmployeeId*"
EXACT_ROLE_HEADER = "Role*"
EXACT_BRANCH_HEADER = "BranchName*"
EXACT_TENANT_HEADER = "Tenant*"
EXACT_MAIN_ID_HEADER = "Id*"

# Guidelines theo chuẩn thiết kế spec
GUIDELINE_EMP_ID = (
    "1. Every entry in this column must also be in one of the columns [Id] in the sheet [Employees]\n"
    "2. NOTE: The contents of this sheet will be merged into any of the sheets [Employees, EmployeeUpdates] "
    "and only migrated when those sheets are migrated. It is not possible to migrate this data by itself."
)
GUIDELINE_ROLE = "Value must match the description of a role listed in Roles and Permissions."
GUIDELINE_BRANCH = "Branch: Must match Branch Name exactly. Empty values will be migrated to HQ."
GUIDELINE_TENANT = (
    'Tenant: Must match url prefix. Example, for demo3.alayacare.com, cell value should be "demo3"'
)

# Danh sách Admin mặc định
DEFAULT_TARGET_USERS = []


def create_employee_roles_with_specific_users(
    workbook: openpyxl.Workbook,
    target_users: Optional[List[Dict[str, Optional[str]]]] = None,
    sheet_name: str = "EmployeeRoles",
    employees_sheet_name: str = "Employees",
) -> openpyxl.Workbook:
    """Tạo worksheet 'EmployeeRoles', ghi các user được chỉ định và định dạng Hàng 2 (Guideline) chữ nghiêng + màu xám."""
    if target_users is None:
        target_users = DEFAULT_TARGET_USERS

    # 1. Nếu sheet đã tồn tại thì xóa đi để làm mới
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]

    sheet = workbook.create_sheet(title=sheet_name)

    # 2. Tạo Hàng 1 (Headers)
    headers = [EXACT_EMP_ID_HEADER, EXACT_ROLE_HEADER, EXACT_BRANCH_HEADER, EXACT_TENANT_HEADER]
    sheet.append(headers)

    # 3. Tạo Hàng 2 (Guidelines)
    guidelines = [GUIDELINE_EMP_ID, GUIDELINE_ROLE, GUIDELINE_BRANCH, GUIDELINE_TENANT]
    sheet.append(guidelines)

    # 🎨 ĐIỀU CHỈNH FONT CHỮ NGHIÊNG VÀ MÀU XÁM CHO HÀNG 2
    # Mã màu '595959' hoặc '7F7F7F' là màu xám chuẩn Excel
    italic_gray_font = Font(name="Calibri", size=10, italic=True, color="595959")
    wrap_alignment = Alignment(wrap_text=True, vertical="top")

    for col_idx in range(1, len(guidelines) + 1):
        cell = sheet.cell(row=2, column=col_idx)
        cell.font = italic_gray_font
        cell.alignment = wrap_alignment  # Giúp xuống dòng đẹp mắt nếu text dài

    # 4. Quét danh sách Employee Id thực tế từ sheet Employees
    existing_emp_ids = set()
    if employees_sheet_name in workbook.sheetnames:
        emp_sheet = workbook[employees_sheet_name]
        id_col_idx = None

        for col in range(1, emp_sheet.max_column + 1):
            cell_val = emp_sheet.cell(row=1, column=col).value
            if cell_val is not None and str(cell_val) == EXACT_MAIN_ID_HEADER:
                id_col_idx = col
                break

        if id_col_idx:
            for r in range(3, emp_sheet.max_row + 1):
                val = emp_sheet.cell(row=r, column=id_col_idx).value
                if val is not None and str(val).strip() != "":
                    existing_emp_ids.add(str(val).strip())

    # 5. Kiểm tra tính tồn tại và ghi các user vào sheet (từ Hàng 3)
    added_count = 0
    skipped_users = []

    for user in target_users:
        emp_id = str(user.get("employee_id", "")).strip()

        if emp_id in existing_emp_ids:
            row_data = [
                emp_id,
                user.get("role", "System Administrator"),
                user.get("branch", None),
                user.get("tenant", "raykay"),
            ]
            sheet.append(row_data)
            added_count += 1
        else:
            skipped_users.append(emp_id)

    print(
        f"✨ Sheet '{sheet_name}': Đã tạo mới, thêm {added_count} user và áp dụng style (In nghiêng + Xám) cho Hàng 2."
    )
    if skipped_users:
        print(f"⚠️ Bỏ qua {len(skipped_users)} user không tồn tại: {skipped_users}")

    return workbook