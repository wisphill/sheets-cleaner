import openpyxl

TARGET_TENANT_HEADERS = ["Tenant*"]
TENANT_VALUE = "raykay"

def update_tenant_column(
    workbook: openpyxl.Workbook, tenant_value: str = TENANT_VALUE
) -> openpyxl.Workbook:
    """Đảm bảo giá trị cột 'Tenant*' là 'raykay' cho tất cả các bản ghi ở dòng 3 trở đi.

    - Kiểm tra tiêu đề ở Hàng 1.
    - Hàng 2 là Guideline -> Bỏ qua.
    - Cập nhật giá trị 'raykay' từ Hàng 3 đến max_row.
    - Nếu không thấy cột ở Hàng 1 -> Bỏ qua worksheet đó.
    """

    for sheet in workbook.worksheets:
        target_col_idx = None

        # 1. Tìm vị trí cột Tenant* ở dòng 1
        for col_idx in range(1, sheet.max_column + 1):
            cell_value = sheet.cell(row=1, column=col_idx).value
            if cell_value is not None:
                clean_value = str(cell_value).strip()
                if clean_value in TARGET_TENANT_HEADERS:
                    target_col_idx = col_idx
                    break

        # 2. Tiến hành cập nhật giá trị từ dòng 3 trở đi
        if target_col_idx:
            updated_count = 0
            for row_idx in range(3, sheet.max_row + 1):
                cell = sheet.cell(row=row_idx, column=target_col_idx)
                # Cập nhật ô thành giá trị tenant_value (vd: "raykay")
                if cell.value != tenant_value:
                    cell.value = tenant_value
                    updated_count += 1

            header_name = sheet.cell(row=1, column=target_col_idx).value
            print(
                f"📝 Sheet '{sheet.title}': Đã gán giá trị '{tenant_value}' cho cột '{header_name}' ({updated_count} dòng được cập nhật)."
            )
        else:
            print(
                f"⏩ Sheet '{sheet.title}': Không tìm thấy cột 'Tenant*' ở dòng 1 -> Bỏ qua."
            )

    return workbook