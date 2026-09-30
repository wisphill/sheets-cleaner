import copy
import openpyxl
from typing import Set

# Khai báo tên Sheet và Cột chuẩn
CONTACTS_SHEET_NAME = "Contacts"
CLIENT_CONTACTS_SHEET_NAME = "ClientContacts"

CONTACT_ID_HEADER = "Id*"
FIRST_NAME_HEADER_KEYWORDS = ["FirstName*", "FirstName", "FirstNames"]
CLIENT_CONTACT_CONTACT_ID_HEADER = "ContactId*"


def is_empty_or_none(val) -> bool:
    """Helper kiểm tra giá trị rỗng."""
    return val is None or str(val).strip() == ""


def remove_dnu_contacts_and_client_contacts(workbook: openpyxl.Workbook) -> openpyxl.Workbook:
    """
    1. Xóa tất cả các hàng trong worksheet 'Contacts' có cột 'FirstName*' chứa 'DNU'.
    2. Thu thập danh sách 'Id*' của các contact này.
    3. Filter worksheet 'ClientContacts' theo cột 'ContactId*' và xóa các hàng liên quan.
    """
    if CONTACTS_SHEET_NAME not in workbook.sheetnames:
        print(f"⏩ Worksheet '{CONTACTS_SHEET_NAME}' không tồn tại trong workbook -> Bỏ qua.")
        return workbook

    sheet_contacts = workbook[CONTACTS_SHEET_NAME]

    # --- BƯỚC 1: XÁC ĐỊNH CỘT TRONG SHEET CONTACTS ---
    col_indices_contacts = {}
    for col in range(1, sheet_contacts.max_column + 1):
        cell_val = sheet_contacts.cell(row=1, column=col).value
        if cell_val is not None:
            col_indices_contacts[str(cell_val).strip()] = col

    id_col_idx = col_indices_contacts.get(CONTACT_ID_HEADER)
    
    # Tìm cột FirstName (hỗ trợ nhiều tên tiêu đề phổ biến)
    first_name_col_idx = None
    for kw in FIRST_NAME_HEADER_KEYWORDS:
        if kw in col_indices_contacts:
            first_name_col_idx = col_indices_contacts[kw]
            break

    if not id_col_idx or not first_name_col_idx:
        print(f"⚠️ Sheet '{CONTACTS_SHEET_NAME}': Không tìm thấy cột '{CONTACT_ID_HEADER}' hoặc 'FirstName' ở Hàng 1 -> Bỏ qua.")
        return workbook

    # --- BƯỚC 2: LỌC DÒNG TRONG CONTACTS VÀ THU THẬP ID BỊ XÓA (DNU) ---
    removed_contact_ids: Set[str] = set()
    total_cols_contacts = sheet_contacts.max_column
    total_rows_contacts = sheet_contacts.max_row

    write_row_contacts = 3  # Giữ lại Hàng 1 (Header) và Hàng 2 (Sub-header / Format gốc)
    deleted_contacts_count = 0

    for read_row in range(3, total_rows_contacts + 1):
        first_name_val = sheet_contacts.cell(row=read_row, column=first_name_col_idx).value
        contact_id_val = sheet_contacts.cell(row=read_row, column=id_col_idx).value
        
        first_name_str = str(first_name_val).strip() if not is_empty_or_none(first_name_val) else ""
        contact_id_str = str(contact_id_val).strip() if not is_empty_or_none(contact_id_val) else ""

        # Kiểm tra xem FirstName có chứa "DNU" không
        if "dnu" in first_name_str.lower():
            if contact_id_str:
                removed_contact_ids.add(contact_id_str)
            deleted_contacts_count += 1
            continue  # Bỏ qua dòng này (xóa)

        # Giữ lại dòng: Shift dữ liệu nếu cần
        if write_row_contacts != read_row:
            for c in range(1, total_cols_contacts + 1):
                src_cell = sheet_contacts.cell(row=read_row, column=c)
                dst_cell = sheet_contacts.cell(row=write_row_contacts, column=c)

                dst_cell.value = src_cell.value
                if src_cell.has_style:
                    dst_cell.number_format = src_cell.number_format
                    dst_cell.font = copy.copy(src_cell.font)
                    dst_cell.border = copy.copy(src_cell.border)
                    dst_cell.fill = copy.copy(src_cell.fill)
                    dst_cell.alignment = copy.copy(src_cell.alignment)

        write_row_contacts += 1

    # Cắt bỏ các dòng thừa ở cuối sheet Contacts
    new_max_row_contacts = write_row_contacts - 1
    if total_rows_contacts > new_max_row_contacts:
        sheet_contacts.delete_rows(new_max_row_contacts + 1, amount=total_rows_contacts - new_max_row_contacts)

    print(f"✂️ Sheet '{CONTACTS_SHEET_NAME}': Đã xóa {deleted_contacts_count} dòng chứa 'DNU'. Thu thập {len(removed_contact_ids)} Contact ID.")

    # --- BƯỚC 3: XÓA CÁC DÒNG LIÊN QUAN TRONG SHEET CLIENTCONTACTS ---
    if CLIENT_CONTACTS_SHEET_NAME not in workbook.sheetnames:
        print(f"⏩ Worksheet '{CLIENT_CONTACTS_SHEET_NAME}' không tồn tại trong workbook -> Bỏ qua phần liên kết.")
        return workbook

    sheet_client_contacts = workbook[CLIENT_CONTACTS_SHEET_NAME]

    col_indices_cc = {}
    for col in range(1, sheet_client_contacts.max_column + 1):
        cell_val = sheet_client_contacts.cell(row=1, column=col).value
        if cell_val is not None:
            col_indices_cc[str(cell_val).strip()] = col

    cc_contact_id_col_idx = col_indices_cc.get(CLIENT_CONTACT_CONTACT_ID_HEADER)

    if not cc_contact_id_col_idx:
        print(f"⚠️ Sheet '{CLIENT_CONTACTS_SHEET_NAME}': Không tìm thấy cột '{CLIENT_CONTACT_CONTACT_ID_HEADER}' -> Bỏ qua.")
        return workbook

    total_cols_cc = sheet_client_contacts.max_column
    total_rows_cc = sheet_client_contacts.max_row

    write_row_cc = 3
    deleted_cc_count = 0

    for read_row in range(3, total_rows_cc + 1):
        cc_contact_id_val = sheet_client_contacts.cell(row=read_row, column=cc_contact_id_col_idx).value
        cc_contact_id_str = str(cc_contact_id_val).strip() if not is_empty_or_none(cc_contact_id_val) else ""

        # Nếu ContactId nằm trong danh sách đã bị xóa ở sheet Contacts -> Bỏ qua (Xóa)
        if cc_contact_id_str in removed_contact_ids:
            deleted_cc_count += 1
            continue

        # Giữ lại dòng: Shift dữ liệu nếu cần
        if write_row_cc != read_row:
            for c in range(1, total_cols_cc + 1):
                src_cell = sheet_client_contacts.cell(row=read_row, column=c)
                dst_cell = sheet_client_contacts.cell(row=write_row_cc, column=c)

                dst_cell.value = src_cell.value
                if src_cell.has_style:
                    dst_cell.number_format = src_cell.number_format
                    dst_cell.font = copy.copy(src_cell.font)
                    dst_cell.border = copy.copy(src_cell.border)
                    dst_cell.fill = copy.copy(src_cell.fill)
                    dst_cell.alignment = copy.copy(src_cell.alignment)

        write_row_cc += 1

    # Cắt bỏ các dòng thừa ở cuối sheet ClientContacts
    new_max_row_cc = write_row_cc - 1
    if total_rows_cc > new_max_row_cc:
        sheet_client_contacts.delete_rows(new_max_row_cc + 1, amount=total_rows_cc - new_max_row_cc)

    print(f"✂️ Sheet '{CLIENT_CONTACTS_SHEET_NAME}': Đã xóa {deleted_cc_count} liên kết ClientContact liên quan.")

    return workbook