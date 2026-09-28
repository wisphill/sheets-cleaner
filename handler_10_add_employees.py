import os
import random
from dotenv import load_dotenv
import openpyxl

# Load biến môi trường từ file .env
load_dotenv()

EXACT_EMPLOYEES_SHEET = "Employees"


def generate_random_test_id() -> str:
    """Sinh chuỗi Id dạng RK_TEST_ + 9 chữ số ngẫu nhiên."""
    random_9_digits = "".join([str(random.randint(0, 9)) for _ in range(9)])
    return f"RK_TEST_{random_9_digits}"


def append_test_user_to_employees(
    workbook: openpyxl.Workbook,
    sheet_name: str = EXACT_EMPLOYEES_SHEET,
    last_name: str = "Nguyen",
    first_name: str = "An",
    email: str = "an@mayflyventures.com",
    env_password_key: str = "TEST_USER_PASSWORD",
) -> openpyxl.Workbook:
    """Append thêm 1 user test vào worksheet Employees với dữ liệu được chỉ
    định."""
    if sheet_name not in workbook.sheetnames:
        print(
            f"⚠️ Worksheet '{sheet_name}' không tồn tại trong workbook -> Bỏ qua."
        )
        return workbook

    sheet = workbook[sheet_name]

    password_val = os.getenv(env_password_key)

    # Sinh Id ngẫu nhiên
    test_id = generate_random_test_id()

    # 1. Bật map tiêu đề cột ở Hàng 1
    header_map = {}
    for col in range(1, sheet.max_column + 1):
        cell_val = sheet.cell(row=1, column=col).value
        if cell_val is not None:
            header_map[str(cell_val).strip()] = col

    # 2. Chuẩn bị danh sách giá trị cho toàn bộ các cột (đảm bảo đúng vị trí cột)
    row_data = [None] * sheet.max_column

    # Gán các giá trị bắt buộc
    if "Id*" in header_map:
        row_data[header_map["Id*"] - 1] = test_id
    if "LastName*" in header_map:
        row_data[header_map["LastName*"] - 1] = last_name
    if "FirstName*" in header_map:
        row_data[header_map["FirstName*"] - 1] = first_name
    if "EmailAddress*" in header_map:
        row_data[header_map["EmailAddress*"] - 1] = email
    if "Password*" in header_map:
        row_data[header_map["Password*"] - 1] = password_val
    if "Tenant*" in header_map:
        row_data[header_map["Tenant*"] - 1] = "raykay"
    if "EmploymentType" in header_map:
        row_data[header_map["EmploymentType"] - 1] = "Casual"
    if "Status*" in header_map:
        row_data[header_map["Status*"] - 1] = "Active"
    if "PayrollId" in header_map:
        row_data[header_map["PayrollId"] - 1] = "*Place-Holder"

    # 3. Append dòng mới vào sheet
    sheet.append(row_data)

    print(
        f"➕ [Employees] Đã thêm thành công user test mới:\n"
        f"   - Id*: {test_id}\n"
        f"   - Name: {last_name} {first_name}\n"
        f"   - Email*: {email}\n"
        f"   - Password*: {'*' * len(password_val)}"
    )

    return workbook