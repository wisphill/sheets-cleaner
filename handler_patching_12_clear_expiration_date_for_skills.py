import openpyxl

EMPLOYEE_SKILLS_SHEET_NAME = "EmployeeSkills"
EXACT_SKILL_HEADER = "Skill*"
EXACT_EXPIRATION_DATE_HEADER = "Expiration Date"

# Skills that must not carry an 'Expiration Date' (no [Expiration Date Label]
# marker in the target system's [Skills] reference sheet).
SKILLS_WITHOUT_EXPIRATION_DATE = {
    "A Palliative Approach To Care",
    "Advanced Dementia Care Training",
    "Arabic",
    "Cantonese",
    "Cert III in Aged Care",
    "Cert III in Aged Care or other",
    "Cert IV in Aged Care or other",
    "Cert IV in Community Services",
    "Croatian",
    "Czech",
    "Experience in Bed Bath",
    "Experience in Complex Care",
    "Experience in Dementia",
    "Experience in Palliative Care",
    "Experience in Personal Care",
    "Female CG",
    "Fijian",
    "Filipino",
    "French",
    "German",
    "Greek",
    "Hungarian",
    "Indonesian",
    "Introduction to Motor Neurone Disease",
    "Introduction to Multiple Sclerosis",
    "Introduction to Parkinson\u2019s Disease",
    "Italian",
    "Korean",
    "Mandarin",
    "Nepali",
    "Polish",
    "Russian",
    "Serbian",
    "Slovak",
    "Spanish",
    "Swedish",
    "Tagalog",
    "Thai",
    "Vaccination - COVID-19 Dose 1",
    "Vaccination - COVID-19 Dose 2",
    "Vaccination - COVID-19 Fully Vaccinated plus Boosters",
    "Vietnamese",
}


def _normalise_skill(value: str) -> str:
    """Lowercase + quy về dấu nháy đơn thẳng để khớp cả 'Parkinson’s' lẫn
    "Parkinson's"."""
    return value.strip().lower().replace("\u2019", "'")


_NORMALISED_TARGET_SKILLS = {
    _normalise_skill(s) for s in SKILLS_WITHOUT_EXPIRATION_DATE
}


def handler_employees_table_empty_expiration_date(
    workbook: openpyxl.Workbook,
    sheet_name: str = EMPLOYEE_SKILLS_SHEET_NAME,
) -> openpyxl.Workbook:
    """Clear the 'Expiration Date' column for every row whose 'Skill*' is one of
    SKILLS_WITHOUT_EXPIRATION_DATE.

    - Headers are on Row 1.
    - Row 2 is the Guideline row -> skipped.
    - Data is cleared from Row 3 onwards.
    """
    if sheet_name not in workbook.sheetnames:
        print(
            f"⏩ Worksheet '{sheet_name}' does not exist in workbook -> Skipped."
        )
        return workbook

    sheet = workbook[sheet_name]

    skill_col_idx = None
    expiration_date_col_idx = None

    # --- STEP 1: Find 'Skill*' and 'Expiration Date' column indexes on Row 1 ---
    for col in range(1, sheet.max_column + 1):
        cell_val = sheet.cell(row=1, column=col).value
        if cell_val is None:
            continue

        clean_val = str(cell_val).strip()
        if clean_val == EXACT_SKILL_HEADER:
            skill_col_idx = col
        elif clean_val == EXACT_EXPIRATION_DATE_HEADER:
            expiration_date_col_idx = col

    if not skill_col_idx or not expiration_date_col_idx:
        missing = []
        if not skill_col_idx:
            missing.append(f"'{EXACT_SKILL_HEADER}'")
        if not expiration_date_col_idx:
            missing.append(f"'{EXACT_EXPIRATION_DATE_HEADER}'")
        print(
            f"⚠️ Sheet '{sheet_name}': Could not find column"
            f" {', '.join(missing)} on Row 1 -> Skipped."
        )
        return workbook

    # --- STEP 2: Clear 'Expiration Date' for the targeted skills ---
    cleared_per_skill = {}

    for row in range(3, sheet.max_row + 1):
        skill_val = sheet.cell(row=row, column=skill_col_idx).value
        if skill_val is None:
            continue

        skill_str = str(skill_val).strip()
        if _normalise_skill(skill_str) not in _NORMALISED_TARGET_SKILLS:
            continue

        date_cell = sheet.cell(row=row, column=expiration_date_col_idx)
        if date_cell.value is not None:
            date_cell.value = None
            cleared_per_skill[skill_str] = (
                cleared_per_skill.get(skill_str, 0) + 1
            )

    total_cleared = sum(cleared_per_skill.values())
    print(
        f"🧹 Sheet '{sheet_name}': Cleared 'Expiration Date' for"
        f" {total_cleared} row(s) across {len(cleared_per_skill)} skill(s)."
    )
    for skill, count in sorted(cleared_per_skill.items()):
        print(f"   • {skill}: {count}")

    return workbook
