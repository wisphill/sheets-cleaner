from typing import Iterable, Set
import openpyxl

EMPLOYEE_SKILLS_SHEET_NAME = "EmployeeSkills"
EXACT_SKILL_HEADER = "Skill*"

# Các biến thể dấu nháy đơn hay gặp khi copy tên Skill từ Excel / Word.
APOSTROPHE_VARIANTS = ("\u2019", "\u2018", "\u02bc", "\u00b4", "`")


def normalize_skill_name(val) -> str:
    """Chuẩn hóa tên Skill để so sánh: bỏ khoảng trắng thừa, không phân biệt
    hoa thường, quy mọi biến thể dấu nháy đơn về dấu nháy ASCII."""
    text = str(val).strip().lower()
    for variant in APOSTROPHE_VARIANTS:
        text = text.replace(variant, "'")
    return text


def normalize_skill_set(skills: Iterable[str]) -> Set[str]:
    """Chuẩn hóa cả một danh sách tên Skill."""
    return {normalize_skill_name(s) for s in skills}


def clear_date_column_for_skills(
    workbook: openpyxl.Workbook,
    date_header: str,
    normalized_skills: Set[str],
    sheet_name: str = EMPLOYEE_SKILLS_SHEET_NAME,
) -> openpyxl.Workbook:
    """Xóa dữ liệu ở cột `date_header` đối với mọi dòng có 'Skill*' nằm trong
    `normalized_skills`.

    - Hàng 1: Header
    - Hàng 2: Guideline -> Bỏ qua
    - Hàng 3+: Xóa dữ liệu
    """
    if sheet_name not in workbook.sheetnames:
        print(
            f"⏩ Worksheet '{sheet_name}' không tồn tại trong workbook -> Bỏ qua."
        )
        return workbook

    sheet = workbook[sheet_name]

    skill_col_idx = None
    date_col_idx = None

    # --- BƯỚC 1: Tìm vị trí cột 'Skill*' và cột ngày ở Hàng 1 ---
    for col in range(1, sheet.max_column + 1):
        cell_val = sheet.cell(row=1, column=col).value
        if cell_val is None:
            continue

        clean_val = str(cell_val).strip()
        if clean_val == EXACT_SKILL_HEADER:
            skill_col_idx = col
        elif clean_val == date_header:
            date_col_idx = col

    if not skill_col_idx or not date_col_idx:
        missing = []
        if not skill_col_idx:
            missing.append(f"'{EXACT_SKILL_HEADER}'")
        if not date_col_idx:
            missing.append(f"'{date_header}'")
        print(
            f"⚠️ Sheet '{sheet_name}': Không tìm thấy cột"
            f" {', '.join(missing)} ở Hàng 1 -> Bỏ qua."
        )
        return workbook

    # --- BƯỚC 2: Xóa dữ liệu ngày cho các Skill mục tiêu ---
    cleared_per_skill = {}

    for row in range(3, sheet.max_row + 1):
        skill_val = sheet.cell(row=row, column=skill_col_idx).value
        if skill_val is None:
            continue

        if normalize_skill_name(skill_val) not in normalized_skills:
            continue

        date_cell = sheet.cell(row=row, column=date_col_idx)
        if date_cell.value is not None:
            date_cell.value = None
            skill_str = str(skill_val).strip()
            cleared_per_skill[skill_str] = (
                cleared_per_skill.get(skill_str, 0) + 1
            )

    total_cleared = sum(cleared_per_skill.values())
    print(
        f"🧹 Sheet '{sheet_name}': Đã xóa '{date_header}' cho"
        f" {total_cleared} dòng thuộc {len(cleared_per_skill)} skill."
    )
    for skill, count in sorted(cleared_per_skill.items()):
        print(f"   • {skill}: {count}")

    return workbook
